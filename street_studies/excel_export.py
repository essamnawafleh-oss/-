"""تصدير ملف Excel بنفس هيكل ومعادلات النسخة الأصلية."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

PRICES_SHEET = "أسعار البنود"
STREETS_SHEET = "الشوارع"

_HEAD_FILL = PatternFill("solid", fgColor="1A5F7A")
_GROUP_FILL = PatternFill("solid", fgColor="D6E9EF")
_INPUT_FILL = PatternFill("solid", fgColor="FFF3CD")
_CALC_FILL = PatternFill("solid", fgColor="E7F4EC")
_HEAD_FONT = Font(bold=True, color="FFFFFF")
_TITLE_FONT = Font(bold=True, size=14, color="0D3D4D")
_THIN = Side(style="thin", color="D0D7DA")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_NUM = "#,##0.00"


def _col(i):
    """فهرس يبدأ من صفر ← حرف العمود."""
    return get_column_letter(i + 1)


def _header_row(ws, row, ncols, height=None):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = _HEAD_FILL
        cell.font = _HEAD_FONT
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if height:
        ws.row_dimensions[row].height = height


def build_workbook(store, streets, region="", tender=""):
    items = store.items
    wb = Workbook()

    # ---- أسعار البنود
    ws = wb.active
    ws.title = PRICES_SHEET
    ws.sheet_view.rightToLeft = True
    ws.append(["أسعار البنود"])
    ws.append(["أدخل أو عدّل سعر الوحدة فقط في العمود E."])
    ws.append([])
    ws.append(["الكود", "المجموعة", "اسم البند", "الوحدة", "سعر الوحدة"])
    for it in items:
        ws.append([it["code"], it["group"], it["name"], it["unit"], float(it.get("price") or 0)])
    ws["A1"].font = _TITLE_FONT
    _header_row(ws, 4, 5, 22)
    for r in range(5, 5 + len(items)):
        for c in range(1, 6):
            ws.cell(row=r, column=c).border = _BORDER
        ws.cell(row=r, column=5).fill = _INPUT_FILL
        ws.cell(row=r, column=5).number_format = _NUM
    for letter, w in zip("ABCDE", (12, 22, 28, 10, 14)):
        ws.column_dimensions[letter].width = w
    ws.freeze_panes = "A5"

    # ---- الشوارع
    ws = wb.create_sheet(STREETS_SHEET)
    ws.sheet_view.rightToLeft = True
    base = ["اسم المواطن", "رقم القطعة", "المنطقة", "طبيعة العطاء", "نوع الخلطة",
            "الطول (م)", "العرض (م)", "مساحة الطريق (م²)", "ملاحظات"]
    nb = len(base)
    headers = list(base)
    for it in items:
        headers += [f"{it['name']}\nالكمية ({it['unit']})", "الكلفة"]
    headers.append("إجمالي كلفة الشارع")
    ncols = len(headers)

    group_row = [""] * ncols
    cur = None
    for i, it in enumerate(items):
        if it["group"] != cur:
            group_row[nb + i * 2] = it["group"]
            cur = it["group"]
    group_row[-1] = "الإجمالي"

    ws.append(["نموذج الدراسات السريعة للشوارع"])
    ws.append([f"تصفية التصدير — المنطقة: {region or 'الكل'} | العطاء: {tender or 'الكل'}"])
    ws.append(group_row)
    ws.append(headers)
    ws["A1"].font = _TITLE_FONT
    for c in range(1, ncols + 1):
        cell = ws.cell(row=3, column=c)
        cell.fill = _GROUP_FILL
        cell.font = Font(bold=True, color="0D3D4D")
    _header_row(ws, 4, ncols, 48)

    for si, s in enumerate(streets):
        rn = si + 5
        row = [s.get("name", ""), s.get("pieceNo") or "", s.get("region") or "", s.get("tenderType") or "",
               "خلطة على خلطة" if s.get("mixType") == "على_خلطة" else "خلطة على بيس كورس",
               s.get("length"), s.get("width"), f"=F{rn}*G{rn}", s.get("notes") or ""]
        cost_refs = []
        q = s.get("quantities") or {}
        for ci, it in enumerate(items):
            qc, cc = _col(nb + ci * 2), _col(nb + 1 + ci * 2)
            v = q.get(it["code"])
            row.append(v if v else None)
            row.append(f'=IF({qc}{rn}="","",{qc}{rn}*INDEX(\'{PRICES_SHEET}\'!$E:$E,'
                       f'MATCH("{it["code"]}",\'{PRICES_SHEET}\'!$A:$A,0)))')
            cost_refs.append(f"{cc}{rn}")
        row.append("=" + "+".join(f"N({r})" for r in cost_refs))
        ws.append(row)
        for c in range(1, ncols + 1):
            cell = ws.cell(row=rn, column=c)
            cell.border = _BORDER
            if c > nb and (c - nb) % 2 == 0 or c == ncols or c == 8:
                cell.number_format = _NUM
                cell.fill = _CALC_FILL
        ws.cell(row=rn, column=ncols).font = Font(bold=True)

    widths = [22, 10, 16, 12, 18, 9, 9, 12, 24] + [12, 12] * len(items) + [16]
    for i, w in enumerate(widths):
        ws.column_dimensions[_col(i)].width = w
    ws.freeze_panes = ws.cell(row=5, column=2)

    # ---- الملخص
    ws = wb.create_sheet("الملخص والتقارير")
    ws.sheet_view.rightToLeft = True
    last = 4 + len(streets)
    tc = _col(ncols - 1)
    ws.append(["الملخص والتقارير"])
    ws.append([f"المنطقة: {region or 'الكل'} | العطاء: {tender or 'الكل'}"])
    ws.append([])
    ws.append(["عدد الشوارع", "", "إجمالي الكلفة", "", "عدد البنود"])
    ws.append([len(streets), "", f"=SUM('{STREETS_SHEET}'!{tc}5:{tc}{last})", "", len(items)])
    ws.append([])
    ws.append(["تفصيل حسب البند"])
    ws.append(["الكود", "المجموعة", "اسم البند", "الوحدة", "إجمالي الكمية", "إجمالي الكلفة"])
    for i, it in enumerate(items):
        qc, cc = _col(nb + i * 2), _col(nb + 1 + i * 2)
        ws.append([it["code"], it["group"], it["name"], it["unit"],
                   f"=SUM('{STREETS_SHEET}'!{qc}5:{qc}{last})",
                   f"=SUM('{STREETS_SHEET}'!{cc}5:{cc}{last})"])
    ws["A1"].font = _TITLE_FONT
    ws["A7"].font = Font(bold=True, size=12, color="0D3D4D")
    for c in (1, 3, 5):
        ws.cell(row=4, column=c).font = Font(bold=True, color="56676E")
        ws.cell(row=5, column=c).font = Font(bold=True, size=13)
    ws["C5"].number_format = _NUM
    _header_row(ws, 8, 6, 22)
    for r in range(9, 9 + len(items)):
        for c in range(1, 7):
            ws.cell(row=r, column=c).border = _BORDER
        ws.cell(row=r, column=5).number_format = _NUM
        ws.cell(row=r, column=6).number_format = _NUM
    for letter, w in zip("ABCDEF", (12, 22, 28, 10, 16, 18)):
        ws.column_dimensions[letter].width = w

    # ---- دليل الاستخدام
    ws = wb.create_sheet("دليل الاستخدام")
    ws.sheet_view.rightToLeft = True
    ws.append(["دليل الاستخدام"])
    ws.append([])
    for row in [
        [1, "اختر المنطقة وطبيعة العطاء (خلطات / تصريف)."],
        [2, "نوع الخلطة: على بيس كورس (عادي) أو على خلطة (تسوية وبيس = صفر)."],
        [3, "التسوية = (عرض+2)×طول | بيس كورس = (عرض+1)×طول | خلطة = عرض×طول"],
        [4, "أضف جدراناً متعددة (مسلح أو دبش) — تُحسب الكميات تلقائياً."],
        [5, "صدّر حسب المنطقة وطبيعة العطاء من صفحة التصدير."],
    ]:
        ws.append(row)
    ws["A1"].font = _TITLE_FONT
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 80

    wb.active = 1
    return wb


def export(store, streets, path, region="", tender=""):
    build_workbook(store, streets, region, tender).save(path)
